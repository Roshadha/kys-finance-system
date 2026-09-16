from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document
from pypdf import PdfReader


ROOT = Path(r"C:\Users\user\Desktop\New folder (11)\source_materials")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def inspect_pdf(path: Path) -> dict:
    reader = PdfReader(str(path))
    samples = []
    total_chars = 0
    for index, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        total_chars += len(text)
        if index < 3 or index >= len(reader.pages) - 2:
            samples.append({"page": index + 1, "chars": len(text), "text": text[:700]})
    return {
        "type": "pdf",
        "pages": len(reader.pages),
        "total_chars": total_chars,
        "samples": samples,
    }


def inspect_docx(path: Path) -> dict:
    doc = Document(str(path))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    table_rows = []
    for table in doc.tables:
        for row in table.rows:
            table_rows.append([cell.text.strip() for cell in row.cells])
    return {
        "type": "docx",
        "paragraphs": len(paragraphs),
        "tables": len(doc.tables),
        "sample_paragraphs": paragraphs[:80],
        "sample_table_rows": table_rows[:40],
    }


result = {}
for path in sorted(ROOT.iterdir()):
    item = {"size": path.stat().st_size, "sha256": sha256(path)}
    try:
        item.update(inspect_pdf(path) if path.suffix.lower() == ".pdf" else inspect_docx(path))
    except Exception as exc:
        item["error"] = f"{type(exc).__name__}: {exc}"
    result[path.name] = item

print(json.dumps(result, ensure_ascii=False, indent=2))
