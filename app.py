from __future__ import annotations

import json
import os
import secrets
import shutil
import sqlite3
import sys
import threading
import time
from datetime import UTC, date, datetime, timedelta
from contextlib import closing
from functools import wraps
from io import BytesIO
from pathlib import Path

import pandas as pd
from flask import (
    Flask,
    abort,
    flash,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from flask_sqlalchemy import SQLAlchemy
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import event, func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from sqlalchemy.schema import CreateIndex, CreateTable
from werkzeug.security import check_password_hash, generate_password_hash


BASE_DIR = Path(__file__).resolve().parent
# Bundled templates/static live inside PyInstaller's resource folder, while
# user data must remain beside the executable across upgrades and restarts.
DATA_ROOT = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else BASE_DIR
db = SQLAlchemy()

TASKS = {
    "income_entry": "Income data entry",
    "expense_entry": "Expenditure data entry",
    "vat_entry": "VAT / tax data entry",
    "view_reports": "View reports",
}
ROLES = {"admin": "Administrator", "data_entry": "Data Entry", "management": "Management"}
SUBTOTAL_ONLY_GROUPS = {
    ("income", "Trade Income"),
    ("income", "Inward Taxes"),
    ("expense", "Tax Expenditure"),
}


def utc_now():
    return datetime.now(UTC).replace(tzinfo=None)


class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    display_name = db.Column(db.String(120), nullable=False)
    department = db.Column(db.String(120), default="")
    password_hash = db.Column(db.String(255), nullable=False)
    role = db.Column(db.String(30), nullable=False, default="data_entry")
    tasks_json = db.Column(db.Text, nullable=False, default="[]")
    active = db.Column(db.Boolean, nullable=False, default=True)
    created_at = db.Column(db.DateTime, nullable=False, default=utc_now)

    @property
    def tasks(self):
        try:
            return json.loads(self.tasks_json)
        except (TypeError, json.JSONDecodeError):
            return []

    def can(self, task):
        return self.role == "admin" or task in self.tasks

    def set_password(self, password):
        self.password_hash = generate_password_hash(password)


# Primary keys double as the visible codes.
# Groups: IG_100, IG_200 ... (income) and EG_100, EG_200 ... (expenditure), +100 each.
# Particulars: IP_0010, IP_0020 ... (income) and EP_0010, EP_0020 ... (expenditure), +10 each.
GROUP_PREFIX = {"income": "IG", "expense": "EG"}
CATEGORY_PREFIX = {"income": "IP", "expense": "EP"}
ID_RULES = {"Group": (GROUP_PREFIX, 100, 0), "Category": (CATEGORY_PREFIX, 10, 4)}


class Group(db.Model):
    id = db.Column(db.String(20), primary_key=True)
    kind = db.Column(db.String(20), nullable=False)
    name = db.Column(db.String(100), nullable=False)
    active = db.Column(db.Boolean, nullable=False, default=True)
    __table_args__ = (db.UniqueConstraint("kind", "name", name="uq_group_kind_name"),)


class Category(db.Model):
    id = db.Column(db.String(20), primary_key=True)
    group_id = db.Column(db.String(20), db.ForeignKey("group.id"), nullable=True, index=True)
    kind = db.Column(db.String(20), nullable=False)
    group_name = db.Column(db.String(100), nullable=False)
    name = db.Column(db.String(180), nullable=False)
    service_line = db.Column(db.String(30), nullable=True)
    entity_name = db.Column(db.String(140), nullable=True)
    active = db.Column(db.Boolean, nullable=False, default=True)


class Vehicle(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    identifier = db.Column(db.String(40), unique=True, nullable=False)
    registration_no = db.Column(db.String(50), nullable=True)
    description = db.Column(db.String(140), nullable=True)
    active = db.Column(db.Boolean, nullable=False, default=True)


class Transaction(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    cash_date = db.Column(db.Date, nullable=False, index=True)
    category_id = db.Column(db.String(20), db.ForeignKey("category.id"), nullable=False, index=True)
    vehicle_id = db.Column(db.Integer, db.ForeignKey("vehicle.id"), nullable=True)
    reference = db.Column(db.String(80), nullable=False)
    counterparty = db.Column(db.String(140), nullable=False)
    description = db.Column(db.String(240), nullable=True)
    subtotal = db.Column(db.Numeric(14, 2), nullable=False, default=0)
    sscl = db.Column(db.Numeric(14, 2), nullable=False, default=0)
    vat = db.Column(db.Numeric(14, 2), nullable=False, default=0)
    total = db.Column(db.Numeric(14, 2), nullable=False, default=0)
    created_by_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    created_at = db.Column(db.DateTime, nullable=False, default=utc_now)
    category = db.relationship("Category")
    vehicle = db.relationship("Vehicle", backref=db.backref("transactions", lazy=True))
    created_by = db.relationship("User")
    adjustments = db.relationship("Adjustment", backref="transaction", lazy=True)

    @property
    def adjustment_total(self):
        return sum(float(item.amount) for item in self.adjustments)

    @property
    def adjusted_total(self):
        return float(self.total) + self.adjustment_total


class Adjustment(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    transaction_id = db.Column(db.Integer, db.ForeignKey("transaction.id"), nullable=False, index=True)
    effective_date = db.Column(db.Date, nullable=False, index=True)
    amount = db.Column(db.Numeric(14, 2), nullable=False)
    reason = db.Column(db.String(240), nullable=False)
    created_by_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    created_at = db.Column(db.DateTime, nullable=False, default=utc_now)
    created_by = db.relationship("User")


class AuditLog(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    happened_at = db.Column(db.DateTime, nullable=False, default=utc_now, index=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=True)
    action = db.Column(db.String(80), nullable=False)
    entity_type = db.Column(db.String(50), nullable=False)
    entity_id = db.Column(db.String(40), nullable=True)
    details = db.Column(db.Text, nullable=True)
    user = db.relationship("User")


def format_id(prefix, number, width):
    return f"{prefix}_{number:0{width}d}"


def highest_id_number(session, model, prefix):
    top = 0
    for (value,) in session.execute(select(model.id)):
        head, _, tail = str(value).partition("_")
        if head == prefix and tail.isdigit():
            top = max(top, int(tail))
    return top


@event.listens_for(Session, "before_flush")
def assign_account_ids(session, flush_context, instances):
    """Give every new group / particular the next code for its type."""
    for model in (Group, Category):
        pending = [obj for obj in session.new if isinstance(obj, model) and not obj.id]
        if not pending:
            continue
        prefixes, step, width = ID_RULES[model.__name__]
        tops = {}
        for obj in pending:
            prefix = prefixes[obj.kind]
            if prefix not in tops:
                with session.no_autoflush:
                    tops[prefix] = highest_id_number(session, model, prefix)
                for other in session.new:
                    if isinstance(other, model) and other.id:
                        head, _, tail = str(other.id).partition("_")
                        if head == prefix and tail.isdigit():
                            tops[prefix] = max(tops[prefix], int(tail))
            tops[prefix] += step
            obj.id = format_id(prefix, tops[prefix], width)


# Order used to number groups of databases created before group IDs existed.
LEGACY_GROUP_ORDER = [
    ("income", "Trade Income"), ("income", "Bank Interest Income"), ("income", "Rental Income"),
    ("income", "Construction Projects Income"), ("expense", "Direct Trade Expenditure"),
    ("expense", "Operational Expenditure"), ("expense", "Maintenance Expenditure"),
    ("expense", "Construction Projects Expenditure"), ("expense", "Investment Projects Expenditure"),
    ("income", "Inward Taxes"), ("expense", "Tax Expenditure"),
]


def numeric_part(value):
    digits = "".join(ch for ch in str(value) if ch.isdigit())
    return int(digits) if digits else 0


def migrate_integer_ids(engine, backup_directory=None):
    """Atomically upgrade legacy IDs, preserving cash records and audit links."""
    if engine.dialect.name != "sqlite":
        return
    db_path = engine.url.database
    if not db_path or db_path == ":memory:" or not Path(db_path).exists():
        return
    con = sqlite3.connect(db_path, isolation_level=None, timeout=30)
    migrated = False
    try:
        # Reserve the write lock before taking a snapshot or reading mapping data.
        # A separate reader can still take a consistent SQLite backup (including WAL).
        con.execute("BEGIN IMMEDIATE")
        cols = {row[1]: (row[2] or "").upper() for row in con.execute("PRAGMA table_info(category)")}
        if not cols or not (cols["id"].startswith("INT") or "code" in cols):
            con.execute("ROLLBACK")
            return
        backup_dir = Path(backup_directory) if backup_directory else DATA_ROOT / "backups"
        backup_dir.mkdir(parents=True, exist_ok=True)
        backup_file = backup_dir / f"pre-code-migration-{datetime.now():%Y%m%d-%H%M%S-%f}.db"
        with closing(sqlite3.connect(db_path)) as source, closing(sqlite3.connect(backup_file)) as snapshot:
            source.backup(snapshot)
            if snapshot.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
                raise RuntimeError("Migration backup failed validation; no accounts were changed.")

        group_rows = con.execute('SELECT id, kind, name, active FROM "group"').fetchall()
        rank = {key: i for i, key in enumerate(LEGACY_GROUP_ORDER)}
        group_rows.sort(key=lambda r: (rank.get((r[1], r[2]), len(rank)), numeric_part(r[0])))
        cat_rows = con.execute(
            "SELECT id, kind, group_name, name, service_line, entity_name, active, code FROM category"
        ).fetchall()
        cat_rows.sort(key=lambda r: numeric_part(r[0]))

        group_ids, old_group_ids, new_groups, counters = {}, {}, [], {}
        def add_group(kind, name, active):
            counters[kind] = counters.get(kind, 0) + 1
            gid = format_id(GROUP_PREFIX[kind], counters[kind] * 100, 0)
            group_ids[(kind, name)] = gid
            new_groups.append((gid, kind, name, active))
            return gid
        for old_id, kind, name, active in group_rows:
            old_group_ids[str(old_id)] = add_group(kind, name, active)
        for _, kind, gname, *_rest in cat_rows:
            if (kind, gname) not in group_ids:
                add_group(kind, gname, 1)

        cat_ids, legacy_codes, new_cats, counters = {}, {}, [], {}
        for old_id, kind, gname, name, service, entity, active, code in cat_rows:
            counters[kind] = counters.get(kind, 0) + 1
            cid = format_id(CATEGORY_PREFIX[kind], counters[kind] * 10, 4)
            cat_ids[str(old_id)] = cid
            legacy_codes[str(old_id)] = code
            new_cats.append((cid, group_ids[(kind, gname)], kind, gname, name, service, entity, active))

        def read_table(name):
            cursor = con.execute(f'SELECT * FROM "{name}"')
            return [column[0] for column in cursor.description], cursor.fetchall()

        txn_cols, txns = read_table("transaction")
        audit_cols, audits = read_table("audit_log")
        category_index = txn_cols.index("category_id")
        converted_txns = []
        for row in txns:
            row = list(row)
            row[category_index] = cat_ids[str(row[category_index])]
            converted_txns.append(row)
        entity_type_index = audit_cols.index("entity_type")
        entity_id_index = audit_cols.index("entity_id")
        converted_audits = []
        for row in audits:
            row = list(row)
            value = row[entity_id_index]
            mapping = {"category": cat_ids, "group": old_group_ids}.get(row[entity_type_index], {})
            row[entity_id_index] = mapping.get(str(value), str(value)) if value is not None else None
            converted_audits.append(row)

        # All DDL and data changes participate in the same transaction. No user,
        # vehicle or adjustment records are rebuilt.
        for model in (Transaction, Category, Group, AuditLog):
            con.execute(f'DROP TABLE "{model.__tablename__}"')
        for model in (Group, Category, Transaction, AuditLog):
            table = model.__table__
            con.execute(str(CreateTable(table).compile(dialect=engine.dialect)))
            for index in table.indexes:
                con.execute(str(CreateIndex(index).compile(dialect=engine.dialect)))
        con.executemany('INSERT INTO "group" (id, kind, name, active) VALUES (?,?,?,?)', new_groups)
        con.executemany(
            "INSERT INTO category (id, group_id, kind, group_name, name, service_line, entity_name, active) "
            "VALUES (?,?,?,?,?,?,?,?)", new_cats)
        for table_name, columns, rows in (
                ("transaction", txn_cols, converted_txns), ("audit_log", audit_cols, converted_audits)):
            marks = ",".join("?" * len(columns))
            names = ",".join(f'"{column}"' for column in columns)
            con.executemany(f'INSERT INTO "{table_name}" ({names}) VALUES ({marks})', rows)

        # Keep old account codes as a historical cross-reference even though the
        # UI now uses automatic IDs and no longer asks for a separate account code.
        con.execute("CREATE TABLE account_code_migration ("
                    "entity_type TEXT NOT NULL, old_id TEXT NOT NULL, new_id TEXT NOT NULL, "
                    "legacy_code TEXT, PRIMARY KEY (entity_type, old_id))")
        con.executemany("INSERT INTO account_code_migration VALUES (?,?,?,?)",
                        [("group", old, new, None) for old, new in old_group_ids.items()] +
                        [("category", old, new, legacy_codes[old]) for old, new in cat_ids.items()])
        if con.execute("PRAGMA foreign_key_check").fetchone() is not None:
            raise RuntimeError("Account migration found a broken data reference; changes were rolled back.")
        if con.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
            raise RuntimeError("Account migration failed integrity checks; changes were rolled back.")
        con.execute("COMMIT")
        migrated = True
    except Exception:
        if con.in_transaction:
            con.execute("ROLLBACK")
        raise
    finally:
        con.close()
    if migrated:
        engine.dispose()


def create_app(test_config=None):
    app = Flask(__name__)
    data_dir = DATA_ROOT / "data"
    data_dir.mkdir(exist_ok=True)
    secret_file = data_dir / ".secret_key"
    if os.environ.get("KYS_SECRET_KEY"):
        secret_key = os.environ["KYS_SECRET_KEY"]
    elif secret_file.exists():
        secret_key = secret_file.read_text(encoding="utf-8").strip()
    else:
        secret_key = secrets.token_hex(32)
        secret_file.write_text(secret_key, encoding="utf-8")
    app.config.update(
        SECRET_KEY=secret_key,
        SQLALCHEMY_DATABASE_URI=os.environ.get(
            "KYS_DATABASE_URL", f"sqlite:///{(DATA_ROOT / 'data' / 'kys_finance.db').as_posix()}"
        ),
        SQLALCHEMY_TRACK_MODIFICATIONS=False,
        MAX_CONTENT_LENGTH=2 * 1024 * 1024,
    )
    if test_config:
        app.config.update(test_config)
    (DATA_ROOT / "backups").mkdir(exist_ok=True)
    db.init_app(app)

    with app.app_context():
        migrate_integer_ids(db.engine, app.config.get("MIGRATION_BACKUP_DIR"))
        db.create_all()
        seed_database()

    register_routes(app)
    register_context(app)

    if not app.config.get("TESTING") and os.environ.get("KYS_DISABLE_BACKUP", "0") != "1":
        start_backup_worker(app)
    return app


def audit(action, entity_type, entity_id=None, details=""):
    db.session.add(AuditLog(
        user_id=session.get("user_id"), action=action, entity_type=entity_type,
        entity_id=entity_id, details=details,
    ))


def current_user():
    uid = session.get("user_id")
    return db.session.get(User, uid) if uid else None


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        user = current_user()
        if not user or not user.active:
            session.clear()
            return redirect(url_for("login", next=request.path))
        return view(*args, **kwargs)
    return wrapped


def role_required(*roles):
    def decorator(view):
        @wraps(view)
        @login_required
        def wrapped(*args, **kwargs):
            if current_user().role not in roles:
                abort(403)
            return view(*args, **kwargs)
        return wrapped
    return decorator


def task_required(task):
    def decorator(view):
        @wraps(view)
        @login_required
        def wrapped(*args, **kwargs):
            if not current_user().can(task):
                abort(403)
            return view(*args, **kwargs)
        return wrapped
    return decorator


def parse_date(value, fallback=None):
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return fallback


def parse_money(value):
    try:
        cleaned = str(value or "0").replace(",", "").strip()
        amount = float(cleaned)
        if not -1_000_000_000_000 < amount < 1_000_000_000_000:
            raise ValueError
        return round(amount, 2)
    except (TypeError, ValueError):
        raise ValueError("Enter a valid amount.")


def period_from_request():
    today = date.today()
    start_default = today.replace(day=1)
    next_month = (start_default.replace(day=28) + timedelta(days=4)).replace(day=1)
    end_default = next_month - timedelta(days=1)
    start = parse_date(request.args.get("start"), start_default)
    end = parse_date(request.args.get("end"), end_default)
    if start > end:
        start, end = end, start
    return start, end


def base_transactions(start, end):
    return (
        Transaction.query.join(Category)
        .filter(Transaction.cash_date.between(start, end))
        .order_by(Transaction.cash_date, Transaction.id)
        .all()
    )


def report_groups():
    # Include historical category groups even when disabled for new entries.
    keys = {(g.kind, g.name) for g in Group.query.all()}
    keys.update((c.kind, c.group_name) for c in Category.query.all())
    return [{"kind": kind, "name": name, "key": json.dumps([kind, name])}
            for kind, name in sorted(keys, key=lambda k: (k[0] != "income", k[1]))]


def selected_report_groups():
    options = report_groups()
    selected = list(dict.fromkeys(request.args.getlist("group")))
    if set(selected) - {g["key"] for g in options}:
        abort(400, description="Unknown report group. Select groups from the report form.")
    return options, selected


def build_report(start, end, report_type="monthly", selected_groups=None):
    selected = {tuple(json.loads(key)) for key in (selected_groups or [])}

    def includes(c):
        if selected and (c.kind, c.group_name) not in selected:
            return False
        if report_type == "direct" and not c.service_line:
            return False
        if report_type == "rental" and not c.entity_name:
            return False
        if report_type == "construction" and "Construction" not in c.group_name:
            return False
        return True

    txns = base_transactions(start, end)
    filtered = []
    for txn in txns:
        c = txn.category
        if not includes(c):
            continue
        filtered.append(txn)

    rows = []
    for txn in filtered:
        c = txn.category
        adjusted = txn.adjusted_total
        rows.append({
            "Category ID": c.id,
            "Date": txn.cash_date.isoformat(),
            "Reference": txn.reference,
            "Type": c.kind.title(),
            "Group": c.group_name,
            "Account Code": c.id,
            "Particular": c.name,
            "Service / Property": c.service_line or c.entity_name or "—",
            "Counterparty": txn.counterparty,
            "Subtotal": float(txn.subtotal),
            "SSCL": float(txn.sscl),
            "VAT": float(txn.vat),
            "Original Total": float(txn.total),
            "Adjustments": txn.adjustment_total,
            "Cash Total": adjusted,
        })

    income = sum(r["Cash Total"] for r in rows if r["Type"] == "Income")
    expense = sum(r["Cash Total"] for r in rows if r["Type"] == "Expense")
    grouped = {}
    for row in rows:
        if report_type == "direct":
            key = row["Service / Property"]
        elif report_type == "rental":
            key = row["Service / Property"]
        else:
            key = row["Group"]
        grouped.setdefault(key, {"income": 0.0, "expense": 0.0})
        grouped[key][row["Type"].lower()] += row["Cash Total"]
    summary = [
        {"label": key, **values, "pl": values["income"] - values["expense"]}
        for key, values in sorted(grouped.items())
    ]
    columns = ["Subtotal", "SSCL", "VAT", "Adjustments", "Cash Total"]
    account_totals = {}
    for row in rows:
        totals = account_totals.setdefault(row["Category ID"], dict.fromkeys(columns, 0.0))
        for col in columns:
            totals[col] += row[col]
    sections = []
    for kind in ("income", "expense"):
        groups = {}
        for c in Category.query.filter_by(kind=kind).order_by(Category.group_name, Category.id).all():
            if not includes(c) or c.id not in account_totals:
                continue
            group = groups.setdefault(c.group_name, {"name": c.group_name, "items": [], "totals": dict.fromkeys(columns, 0.0)})
            amounts = account_totals.get(c.id, dict.fromkeys(columns, 0.0))
            group["items"].append({"code": c.id, "name": c.name, "amounts": amounts})
            for col in columns:
                group["totals"][col] += amounts[col]
        sections.append({"kind": kind, "label": "Income" if kind == "income" else "Expenditure",
                         "groups": list(groups.values()),
                         "totals": {col: sum(g["totals"][col] for g in groups.values()) for col in columns}})
    return {"rows": rows, "summary": summary, "sections": sections, "columns": columns,
            "income": income, "expense": expense, "pl": income - expense}


def export_report_xlsx(start, end, report_type, selected_groups=None):
    data = build_report(start, end, report_type, selected_groups)
    output = BytesIO()
    title_map = {
        "monthly": "Monthly Cash-Based P&L",
        "short": "Short Summary P&L",
        "direct": "Direct Trade P&L",
        "rental": "Rented Building P&L",
        "construction": "Construction Projects P&L",
    }
    summary_rows = [{
        "Particular": item["label"], "Income": item["income"],
        "Expenditure": item["expense"], "Liquid P/L": item["pl"],
    } for item in data["summary"]]
    summary_rows.append({
        "Particular": "TOTAL", "Income": data["income"],
        "Expenditure": data["expense"], "Liquid P/L": data["pl"],
    })
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        grouped_rows = []
        styled_rows = []
        def statement_row(label, amounts=None, code="", bold=False, number=""):
            if bold:
                styled_rows.append(len(grouped_rows) + 6)
            if report_type == "short":
                grouped_rows.append({"Group": label, "Cash Total": amounts["Cash Total"] if amounts else None})
            else:
                amount = amounts["Cash Total"] if amounts else None
                grouped_rows.append({"No.": number, "Particular": label, "A/C code": code,
                                     "Amount (LKR)": amount if not bold else None,
                                     "Total (LKR)": amount if bold else None})
        for section in data["sections"]:
            if not section["groups"]:
                continue
            statement_row(section["label"] + " Particulars", bold=True)
            for group in section["groups"]:
                if report_type == "short":
                    statement_row(group["name"], group["totals"], bold=True)
                else:
                    statement_row(group["name"], bold=True)
                    for number, item in enumerate(group["items"], start=1):
                        statement_row(item["name"], item["amounts"], item["code"], number=number)
                    statement_row(group["name"] + " — Total", group["totals"], bold=True)
            statement_row("Total " + section["label"], section["totals"], bold=True)
        statement_row("Liquid P/L", {col: data["pl"] if col == "Cash Total" else None for col in data["columns"]}, bold=True)
        pd.DataFrame(grouped_rows).to_excel(writer, sheet_name="Group Report", index=False, startrow=4)
        pd.DataFrame(summary_rows).to_excel(writer, sheet_name="Summary", index=False, startrow=4)
        if report_type != "short":
            pd.DataFrame([{k: v for k, v in row.items() if k != "Category ID"} for row in data["rows"]]).to_excel(writer, sheet_name="Detailed Ledger", index=False, startrow=4)
        for ws in writer.book.worksheets:
            ws["A1"] = "K.Y.S. Cash-Based P&L Management System"
            ws["A2"] = title_map.get(report_type, title_map["monthly"])
            ws["A3"] = f"Cash period: {start:%d %b %Y} to {end:%d %b %Y}"
            ws["A4"] = "Groups: " + (", ".join(f"{json.loads(key)[0].title()}: {json.loads(key)[1]}" for key in selected_groups) if selected_groups else "All groups")
            ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
            ws["A1"].fill = PatternFill("solid", fgColor="0D3B34")
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(ws.max_column, 2 if report_type == "short" else 4))
            ws["A2"].font = Font(size=12, bold=True, color="163C34")
            header_fill = PatternFill("solid", fgColor="DCEBE6")
            thin = Side(style="thin", color="D6E0DC")
            for cell in ws[5]:
                cell.font = Font(bold=True, color="163C34")
                cell.fill = header_fill
                cell.border = Border(bottom=thin)
                cell.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A6"
            ws.auto_filter.ref = f"A5:{ws.cell(5, ws.max_column).column_letter}{ws.max_row}"
            for column_index, column in enumerate(ws.iter_cols(), start=1):
                letter = get_column_letter(column_index)
                max_len = max(len(str(cell.value or "")) for cell in column)
                ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 42)
            for row in ws.iter_rows(min_row=6):
                for cell in row:
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = '#,##0.00;[Red]-#,##0.00'
                    elif cell.data_type == "f":
                        cell.data_type = "s"
        ws = writer.book["Group Report"]
        ws.auto_filter.ref = None
        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 58
        for col in range(3, 6):
            ws.column_dimensions[get_column_letter(col)].width = 20
        for row in ws.iter_rows(min_row=6):
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True,
                                           horizontal="right" if isinstance(cell.value, (int, float)) else "left")
                cell.border = Border(bottom=Side(style="thin", color="D6E0DC"))
            if report_type != "short":
                row[0].number_format = '0'
            ws.row_dimensions[row[0].row].height = 32
        for row_number in styled_rows:
            for cell in ws[row_number]:
                cell.font = Font(bold=True, color="163C34")
                cell.fill = PatternFill("solid", fgColor="EAF2EF")
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:5"
        ws.print_area = f"A1:E{ws.max_row}"
        ws.cell(ws.max_row + 2, 1, "Total Liquid Income − Total Expenditure = Liquid P/L")
        ws.cell(ws.max_row + 1, 1, f"{data['income']:,.2f} − {data['expense']:,.2f} = {data['pl']:,.2f} LKR")
        ws.print_area = f"A1:E{ws.max_row}"
        if report_type == "short":
            ws.column_dimensions["A"].width = 58
            ws.column_dimensions["B"].width = 22
            ws.print_area = f"A1:B{ws.max_row}"
        ws.print_options.horizontalCentered = True
    output.seek(0)
    return output


def register_context(app):
    @app.context_processor
    def shared():
        return {
            "current_user": current_user(), "TASKS": TASKS, "ROLES": ROLES,
            "today": date.today(), "csrf_token": session.get("csrf_token", ""),
        }

    @app.template_filter("money")
    def money(value):
        return f"{float(value or 0):,.2f}"

    @app.before_request
    def csrf_and_session():
        if "csrf_token" not in session:
            session["csrf_token"] = secrets.token_urlsafe(24)
        if request.method == "POST":
            token = request.form.get("csrf_token") or request.headers.get("X-CSRF-Token")
            if not secrets.compare_digest(token or "", session.get("csrf_token", "")):
                abort(400, "Invalid or expired form token. Refresh the page and try again.")


def register_routes(app):
    @app.get("/health")
    def health():
        return {"status": "ok", "database": "connected", "service": "kys-finance"}

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if request.method == "POST":
            username = request.form.get("username", "").strip()
            user = User.query.filter(func.lower(User.username) == username.lower()).first()
            if user and user.active and check_password_hash(user.password_hash, request.form.get("password", "")):
                session.clear()
                session["user_id"] = user.id
                session["csrf_token"] = secrets.token_urlsafe(24)
                audit("login", "user", user.id, "Successful sign in")
                db.session.commit()
                return redirect(request.args.get("next") or url_for("dashboard"))
            flash("The username or password is incorrect.", "error")
        return render_template("login.html")

    @app.post("/logout")
    @login_required
    def logout():
        uid = current_user().id
        audit("logout", "user", uid, "Signed out")
        db.session.commit()
        session.clear()
        return redirect(url_for("login"))

    @app.get("/")
    @login_required
    def dashboard():
        today = date.today()
        start = today.replace(day=1)
        data = build_report(start, today, "monthly")
        recent = Transaction.query.order_by(Transaction.created_at.desc()).limit(7).all()
        monthly = []
        for offset in range(5, -1, -1):
            month = (today.replace(day=1) - timedelta(days=offset * 28)).replace(day=1)
            nxt = (month.replace(day=28) + timedelta(days=4)).replace(day=1)
            report = build_report(month, nxt - timedelta(days=1), "monthly")
            monthly.append({"label": month.strftime("%b"), **report})
        return render_template("dashboard.html", data=data, recent=recent, monthly=monthly)

    @app.get("/transactions")
    @login_required
    def transactions():
        start, end = period_from_request()
        kind = request.args.get("kind", "")
        query = Transaction.query.join(Category).filter(Transaction.cash_date.between(start, end))
        if kind in {"income", "expense"}:
            query = query.filter(Category.kind == kind)
        items = query.order_by(Transaction.cash_date.desc(), Transaction.id.desc()).all()
        return render_template("transactions.html", items=items, start=start, end=end, kind=kind)

    @app.route("/transactions/new", methods=["GET", "POST"])
    @login_required
    def transaction_new():
        requested_kind = request.args.get("kind", request.form.get("kind", "income"))
        if requested_kind not in {"income", "expense"}:
            requested_kind = "income"
        task = "income_entry" if requested_kind == "income" else "expense_entry"
        if not current_user().can(task):
            abort(403)
        categories = Category.query.filter_by(kind=requested_kind, active=True).order_by(Category.group_name, Category.name).all()
        vehicles = Vehicle.query.filter_by(active=True).order_by(Vehicle.identifier).all()
        if request.method == "POST":
            try:
                category = db.session.get(Category, request.form.get("category_id", "").strip())
                selected_group = request.form.get("group_name", "").strip()
                if not selected_group:
                    raise ValueError("Select a group.")
                if not category or category.kind != requested_kind or not category.active or category.group_name != selected_group:
                    raise ValueError("Select a valid account.")
                cash_date = parse_date(request.form.get("cash_date"))
                if not cash_date:
                    raise ValueError("Select the actual cash date.")
                subtotal = parse_money(request.form.get("subtotal"))
                sscl = parse_money(request.form.get("sscl"))
                vat = parse_money(request.form.get("vat"))
                if (requested_kind, selected_group) in SUBTOTAL_ONLY_GROUPS and (sscl or vat):
                    raise ValueError("Enter SSCL and VAT as separate tax particulars, not in these amount fields.")
                if min(subtotal, sscl, vat) < 0:
                    raise ValueError("Amounts cannot be negative.")
                total = round(subtotal + sscl + vat, 2)
                if total <= 0:
                    raise ValueError("Grand total must be greater than zero.")
                reference = request.form.get("reference", "").strip()
                counterparty = request.form.get("counterparty", "").strip()
                vehicle_id = request.form.get("vehicle_id") or None
                if "Vehicle Running" in category.name and not vehicle_id:
                    raise ValueError("Choose the vehicle for this maintenance entry.")
                item = Transaction(
                    cash_date=cash_date, category_id=category.id, vehicle_id=int(vehicle_id) if vehicle_id else None,
                    reference=reference, counterparty=counterparty,
                    description=request.form.get("description", "").strip(), subtotal=subtotal,
                    sscl=sscl, vat=vat, total=total, created_by_id=current_user().id,
                )
                db.session.add(item)
                db.session.flush()
                audit("create", "transaction", item.id, f"{requested_kind} {total:.2f} / {reference}")
                db.session.commit()
                flash(f"Cash {requested_kind} saved. The ledger entry is now locked.", "success")
                return redirect(url_for("transactions"))
            except (ValueError, TypeError) as exc:
                db.session.rollback()
                flash(str(exc), "error")
        subtotal_only_groups = {name for kind, name in SUBTOTAL_ONLY_GROUPS if kind == requested_kind}
        return render_template("transaction_form.html", kind=requested_kind, categories=categories,
                               vehicles=vehicles, subtotal_only_groups=subtotal_only_groups)

    @app.route("/transactions/<int:transaction_id>/adjust", methods=["GET", "POST"])
    @role_required("admin")
    def adjustment_new(transaction_id):
        txn = db.get_or_404(Transaction, transaction_id)
        if request.method == "POST":
            try:
                amount = parse_money(request.form.get("amount"))
                reason = request.form.get("reason", "").strip()
                effective_date = parse_date(request.form.get("effective_date"))
                if amount == 0 or not reason or not effective_date:
                    raise ValueError("Date, a non-zero adjustment, and a reason are required.")
                adj = Adjustment(
                    transaction_id=txn.id, effective_date=effective_date, amount=amount,
                    reason=reason, created_by_id=current_user().id,
                )
                db.session.add(adj)
                db.session.flush()
                audit("adjust", "transaction", txn.id, f"Adjustment {amount:.2f}: {reason}")
                db.session.commit()
                flash("Adjustment posted. The original entry remains unchanged.", "success")
                return redirect(url_for("transactions"))
            except ValueError as exc:
                db.session.rollback()
                flash(str(exc), "error")
        return render_template("adjustment_form.html", txn=txn)

    @app.get("/reports")
    @task_required("view_reports")
    def reports():
        start, end = period_from_request()
        report_type = request.args.get("type", "monthly")
        if report_type not in {"monthly", "short"}:
            report_type = "monthly"
        groups, selected_groups = selected_report_groups()
        data = build_report(start, end, report_type, selected_groups)
        return render_template("reports.html", data=data, start=start, end=end, report_type=report_type,
                               groups=groups, selected_groups=selected_groups)

    @app.get("/reports/export")
    @task_required("view_reports")
    def reports_export():
        start, end = period_from_request()
        report_type = request.args.get("type", "monthly")
        if report_type not in {"monthly", "short"}:
            report_type = "monthly"
        groups, selected_groups = selected_report_groups()
        output = export_report_xlsx(start, end, report_type, selected_groups)
        audit("export", "report", None, f"{report_type}: {start} to {end}; groups: {selected_groups or 'All'}")
        db.session.commit()
        return send_file(
            output, as_attachment=True,
            download_name=f"KYS_{report_type}_PL_{start}_{end}.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @app.get("/history")
    @role_required("admin")
    def history():
        logs = AuditLog.query.order_by(AuditLog.happened_at.desc()).limit(300).all()
        return render_template("history.html", logs=logs)

    @app.get("/users")
    @role_required("admin")
    def users():
        return render_template("users.html", users=User.query.order_by(User.active.desc(), User.display_name).all())

    @app.route("/users/new", methods=["GET", "POST"])
    @role_required("admin")
    def user_new():
        if request.method == "POST":
            try:
                username = request.form.get("username", "").strip()
                password = request.form.get("password", "")
                role = request.form.get("role", "data_entry")
                if not username or len(password) < 8 or role not in ROLES:
                    raise ValueError("Use a username, valid role, and password of at least 8 characters.")
                if User.query.filter(func.lower(User.username) == username.lower()).first():
                    raise ValueError("That username is already in use.")
                user = User(
                    username=username,
                    display_name=request.form.get("display_name", "").strip() or username,
                    department=request.form.get("department", "").strip(), role=role,
                    tasks_json=json.dumps(request.form.getlist("tasks")),
                )
                user.set_password(password)
                db.session.add(user)
                db.session.flush()
                audit("create", "user", user.id, f"Created {username} as {role}")
                db.session.commit()
                flash("User account created.", "success")
                return redirect(url_for("users"))
            except ValueError as exc:
                db.session.rollback()
                flash(str(exc), "error")
        return render_template("user_form.html")

    @app.route("/users/<int:user_id>/edit", methods=["GET", "POST"])
    @role_required("admin")
    def user_edit(user_id):
        user = db.get_or_404(User, user_id)
        if request.method == "POST":
            role = request.form.get("role", "data_entry")
            password = request.form.get("password", "")
            if role not in ROLES:
                flash("Select a valid role.", "error")
            elif password and len(password) < 8:
                flash("A reset password must be at least 8 characters.", "error")
            else:
                user.display_name = request.form.get("display_name", "").strip() or user.username
                user.department = request.form.get("department", "").strip()
                user.role = role
                user.tasks_json = json.dumps(request.form.getlist("tasks"))
                if password:
                    user.set_password(password)
                audit("update", "user", user.id, f"Role/tasks updated for {user.username}")
                db.session.commit()
                flash("User role and task assignments updated.", "success")
                return redirect(url_for("users"))
        return render_template("user_form.html", edit_user=user)

    @app.post("/users/<int:user_id>/toggle")
    @role_required("admin")
    def user_toggle(user_id):
        user = db.get_or_404(User, user_id)
        if user.id == current_user().id:
            flash("You cannot disable your own signed-in account.", "error")
        else:
            user.active = not user.active
            audit("status", "user", user.id, f"Active: {user.active}")
            db.session.commit()
            flash("User access updated.", "success")
        return redirect(url_for("users"))

    @app.get("/vehicles")
    @role_required("admin")
    def vehicles():
        return render_template("vehicles.html", vehicles=Vehicle.query.order_by(Vehicle.identifier).all())

    @app.get("/categories")
    @role_required("admin")
    def categories():
        items = Category.query.order_by(Category.kind, Category.id).all()
        groups = Group.query.filter_by(active=True).order_by(Group.kind, Group.id).all()
        return render_template("categories.html", categories=items, groups=groups)

    @app.post("/categories")
    @role_required("admin")
    def category_new():
        kind = request.form.get("kind", "")
        group_name = request.form.get("group_name", "").strip()
        name = request.form.get("name", "").strip()
        valid_group = kind in {"income", "expense"} and Group.query.filter_by(
            kind=kind, name=group_name, active=True
        ).first()
        if kind not in {"income", "expense"} or not group_name or not name:
            flash("Type, group, and particular name are required.", "error")
        elif not valid_group:
            flash("Select a valid group for the chosen type, or add a new one on the Groups page first.", "error")
        else:
            item = Category(
                group_id=valid_group.id,
                kind=kind,
                group_name=group_name, name=name,
                service_line=request.form.get("service_line", "").strip() or None,
                entity_name=request.form.get("entity_name", "").strip() or None,
            )
            try:
                db.session.add(item)
                db.session.flush()  # the next IP_ id is generated here
                audit("create", "category", item.id, f"{kind}: {name}")
                db.session.commit()
                flash(f"Account particular added with code {item.id}.", "success")
            except IntegrityError:
                db.session.rollback()
                flash("Could not generate a unique ID. Please try again.", "error")
        return redirect(url_for("categories"))

    @app.post("/categories/<category_id>/toggle")
    @role_required("admin")
    def category_toggle(category_id):
        item = db.get_or_404(Category, category_id)
        item.active = not item.active
        audit("status", "category", item.id, f"Active: {item.active}")
        db.session.commit()
        flash("Account availability updated.", "success")
        return redirect(url_for("categories"))

    @app.get("/groups")
    @role_required("admin")
    def groups():
        items = Group.query.order_by(Group.kind, Group.name).all()
        return render_template("groups.html", groups=items)

    @app.post("/groups")
    @role_required("admin")
    def group_new():
        kind = request.form.get("kind", "")
        name = request.form.get("name", "").strip()
        if kind not in {"income", "expense"} or not name:
            flash("Type and group name are required.", "error")
        elif Group.query.filter(Group.kind == kind, func.lower(Group.name) == name.lower()).first():
            flash("That group already exists for this type.", "error")
        else:
            item = Group(kind=kind, name=name)
            try:
                db.session.add(item)
                db.session.flush()  # the next IG_ id is generated here
                audit("create", "group", item.id, f"{kind}: {name}")
                db.session.commit()
                flash(f"Group added with code {item.id}.", "success")
            except IntegrityError:
                db.session.rollback()
                flash("Could not generate a unique ID. Please try again.", "error")
        return redirect(url_for("groups"))

    @app.post("/groups/<group_id>/toggle")
    @role_required("admin")
    def group_toggle(group_id):
        item = db.get_or_404(Group, group_id)
        item.active = not item.active
        audit("status", "group", item.id, f"Active: {item.active}")
        db.session.commit()
        flash("Group availability updated.", "success")
        return redirect(url_for("groups"))

    @app.post("/vehicles")
    @role_required("admin")
    def vehicle_new():
        identifier = request.form.get("identifier", "").strip()
        if not identifier:
            flash("Vehicle identifier is required.", "error")
        elif Vehicle.query.filter(func.lower(Vehicle.identifier) == identifier.lower()).first():
            flash("That vehicle identifier already exists.", "error")
        else:
            vehicle = Vehicle(
                identifier=identifier,
                registration_no=request.form.get("registration_no", "").strip(),
                description=request.form.get("description", "").strip(),
            )
            db.session.add(vehicle)
            db.session.flush()
            audit("create", "vehicle", vehicle.id, identifier)
            db.session.commit()
            flash("Vehicle added.", "success")
        return redirect(url_for("vehicles"))

    @app.post("/backup")
    @role_required("admin")
    def backup_now():
        path, external_path, external_error = make_backup(app, force_external=True)
        audit("backup", "database", None, path.name)
        db.session.commit()
        if external_error:
            flash(f"Local backup created, but the external backup failed: {external_error}", "error")
        elif external_path:
            flash(f"Backup created locally and at {external_path}", "success")
        else:
            flash(f"Local backup created: {path.name}. Configure an external backup for disaster recovery.", "success")
        return redirect(url_for("dashboard"))

    @app.route("/account/password", methods=["GET", "POST"])
    @login_required
    def change_password():
        user = current_user()
        if request.method == "POST":
            current = request.form.get("current_password", "")
            new = request.form.get("new_password", "")
            if not check_password_hash(user.password_hash, current):
                flash("Current password is incorrect.", "error")
            elif len(new) < 8:
                flash("New password must be at least 8 characters.", "error")
            else:
                user.set_password(new)
                audit("password", "user", user.id, "Password changed")
                db.session.commit()
                flash("Password updated.", "success")
                return redirect(url_for("dashboard"))
        return render_template("change_password.html")

    @app.errorhandler(403)
    def forbidden(_):
        return render_template("error.html", code=403, message="You do not have permission to open this area."), 403

    @app.errorhandler(404)
    def not_found(_):
        return render_template("error.html", code=404, message="The requested record or page was not found."), 404


def make_backup(app, force_external=False):
    database_uri = app.config["SQLALCHEMY_DATABASE_URI"]
    if not database_uri.startswith("sqlite:///"):
        raise RuntimeError("Automatic file backup is only available for SQLite.")
    source = Path(database_uri.removeprefix("sqlite:///"))
    if not source.is_absolute():
        source = DATA_ROOT / source
    destination = DATA_ROOT / "backups" / f"kys_finance_{datetime.now():%Y%m%d_%H%M%S_%f}.db"
    temporary = destination.with_suffix(".tmp")
    destination.parent.mkdir(exist_ok=True)
    with closing(sqlite3.connect(source)) as source_db, closing(sqlite3.connect(temporary)) as backup_db:
        source_db.backup(backup_db)
        if backup_db.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
            raise RuntimeError("SQLite backup integrity check failed.")
    temporary.replace(destination)
    backups = sorted(destination.parent.glob("kys_finance_*.db"), key=lambda p: p.stat().st_mtime, reverse=True)
    for old in backups[60:]:
        old.unlink(missing_ok=True)
    external_path = None
    external_error = None
    if (DATA_ROOT / "backup-target.txt").exists():
        try:
            external_path = copy_external_backup(destination, force=force_external)
        except (OSError, RuntimeError, sqlite3.DatabaseError) as exc:
            external_error = str(exc)
            app.logger.exception("External backup failed; local backup remains available")
    return destination, external_path, external_error


def copy_external_backup(local_backup, force=False):
    target = Path((DATA_ROOT / "backup-target.txt").read_text(encoding="utf-8-sig").strip())
    if not target.is_absolute() or not target.is_dir():
        raise RuntimeError("Configured external backup folder is unavailable.")
    if target.resolve() == DATA_ROOT.resolve() or DATA_ROOT.resolve() in target.resolve().parents:
        raise RuntimeError("External backup folder must be outside the application folder.")
    day = datetime.now().strftime("%Y-%m-%d")
    stamp = datetime.now().strftime("%H%M%S_%f")
    backup_dir = target / "KYS-Finance-Backups" / (f"{day}_{stamp}" if force else day)
    complete = backup_dir / "backup-complete.txt"
    if (not force and complete.exists() and (backup_dir / "kys_finance.db").exists()
            and (backup_dir / ".secret_key").exists()):
        return backup_dir
    backup_dir.mkdir(parents=True, exist_ok=True)
    temporary = backup_dir / "kys_finance.db.tmp"
    shutil.copy2(local_backup, temporary)
    with closing(sqlite3.connect(temporary)) as check_db:
        if check_db.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
            raise RuntimeError("External backup integrity check failed.")
    temporary.replace(backup_dir / "kys_finance.db")
    secret = DATA_ROOT / "data" / ".secret_key"
    secret_temp = backup_dir / ".secret_key.tmp"
    if secret.exists():
        shutil.copy2(secret, secret_temp)
    elif os.environ.get("KYS_SECRET_KEY"):
        secret_temp.write_text(os.environ["KYS_SECRET_KEY"], encoding="utf-8")
    else:
        raise RuntimeError("Session secret is unavailable for the external backup.")
    secret_temp.replace(backup_dir / ".secret_key")
    complete.write_text(f"Verified {datetime.now().isoformat(timespec='seconds')}\n", encoding="utf-8")
    return backup_dir


def start_backup_worker(app):
    if getattr(app, "_backup_started", False):
        return
    app._backup_started = True

    def worker():
        while True:
            try:
                make_backup(app)
            except Exception:
                app.logger.exception("Scheduled backup failed")
            time.sleep(2 * 60 * 60)

    threading.Thread(target=worker, daemon=True, name="sqlite-backup").start()


def get_or_create_group(kind, name):
    group = Group.query.filter_by(kind=kind, name=name).first()
    if not group:
        group = Group(kind=kind, name=name)
        db.session.add(group)
        db.session.flush()
    return group


def create_category(kind, group_name, name, service=None, entity=None):
    group = get_or_create_group(kind, group_name)
    item = Category(group_id=group.id, kind=kind, group_name=group_name, name=name,
                    service_line=service, entity_name=entity)
    db.session.add(item)
    db.session.flush()
    return item


def seed_database():
    if not User.query.first():
        defaults = [
            ("admin", "System Administrator", "Accounts", "admin", [*TASKS], "Admin@2026"),
            ("management", "Management", "Management", "management", ["view_reports"], "Manage@2026"),
            ("recovery", "Recovery Department", "Recovery", "data_entry", ["income_entry", "vat_entry"], "KysRec@2026"),
        ]
        for username, display, department, role, tasks, password in defaults:
            user = User(username=username, display_name=display, department=department, role=role, tasks_json=json.dumps(tasks))
            user.set_password(password)
            db.session.add(user)

    if not Category.query.first():
        def add(kind, group, name, service=None, entity=None):
            create_category(kind, group, name, service=service, entity=entity)

        add("income", "Trade Income", "Security Service Income", "Security")
        add("income", "Trade Income", "Cleaning Service Income", "Cleaning")
        add("income", "Bank Interest Income", "Bank Interest Received")
        properties = [
            "K.Y.S. Auto Service", "Pinarawa Housing Complex", "Welekade Shops", "Ratwatta Guest House",
            "2nd Mile Post Shop Complex", "Chandrasiri House", "MD Home Upper Stairs", "MD Home Shop",
            "Haputhale Hotel", "2nd Mile Post Hostel", "Galaxy - A Hostel", "Galaxy - E Hostel",
            "Pihillakade Housing Complex", "Vineethagama House",
        ]
        for entity in properties:
            add("income", "Rental Income", f"{entity} Rent Income", entity=entity)
        add("income", "Construction Projects Income", "Construction Project Receipt")

        for name, service in [
            ("Security Service Salary", "Security"), ("Security Service EPF 12%", "Security"),
            ("Security Service ETF 3%", "Security"), ("Cleaning Service Salary", "Cleaning"),
            ("Cleaning Service EPF 12%", "Cleaning"), ("Cleaning Service ETF 3%", "Cleaning"),
        ]:
            add("expense", "Direct Trade Expenditure", name, service)
        operational = [
            "Staff Salary Expenditure", "EPF (12%) Expenditure", "ETF (3%) Expenditure", "Office Expenditure",
            "Printing & Stationery Expenditure", "Postage & Courier Expenditure",
            "Travelling & Subsistence Expenditure", "Service Operation Expenditure",
            "Vehicle Running & Maintenance Expenditure", "Security Items Expenditure",
            "Cleaning Items & Materials Expenditure", "Training, Permit, Licence & Miscellaneous Expenditure",
            "Tender Documents Expenditure", "Charity & Donation Expenditure",
            "Marketing, Promotion, Subscription & Periodicals Expenditure", "Professional Fees Expenditure",
            "Staff Training & Programme Expenditure", "Cleaning Service Other Expenditure",
            "Electricity Charges Expenditure", "Water Supply Charges Expenditure",
            "Telephone & Internet Charges Expenditure", "MD House Expenditure",
            "Kottawa Prime Residence Expenditure", "Personal Insurance Premium Expenditure",
            "Employees Insurance Premium Expenditure", "Vehicle & General Insurance Premium Expenditure",
            "Financial / Loan & OD Interest Expenditure",
        ]
        for name in operational:
            add("expense", "Operational Expenditure", name)
        for entity in properties:
            add("expense", "Maintenance Expenditure", f"{entity} Maintenance & Expenditure", entity=entity)
        add("expense", "Construction Projects Expenditure", "Construction Project Payment")
        investments = [
            "WIP Velaudam Land Building Investment - Phase I", "WIP Velaudam Land Building Investment - Phase II",
            "WIP Damanwara St. Bernard Tea Plantation Investment",
            "WIP Damanwara St. Bernard Land Development Investment",
            "WIP Galaxy A & E Hostel Building Investment", "WIP Pihillakade Housing Complex Investment",
            "WIP Kandana Pepper Cultivation Investment",
        ]
        for name in investments:
            add("expense", "Investment Projects Expenditure", name)

    # Link any particular that predates group IDs to its group.
    for cat in Category.query.filter(Category.group_id.is_(None)).all():
        cat.group_id = get_or_create_group(cat.kind, cat.group_name).id

    # Apply new tax accounts to both fresh and existing company databases.
    for kind, group_name in (("income", "Inward Taxes"), ("expense", "Tax Expenditure")):
        get_or_create_group(kind, group_name)
        for particular in ("SSCL", "VAT"):
            if not Category.query.filter_by(kind=kind, group_name=group_name, name=particular).first():
                create_category(kind, group_name, particular)

    if not Vehicle.query.first():
        db.session.add_all([
            Vehicle(identifier="A", description="Vehicle A"),
            Vehicle(identifier="B", description="Vehicle B"),
            Vehicle(identifier="C", description="Vehicle C"),
        ])
    db.session.commit()


app = create_app()


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "5000")), debug=False)
