from pathlib import Path

from docx import Document


ROOT = Path(r"C:\Users\user\Desktop\New folder (11)")
SOURCE = ROOT / "output" / "Grade_10_Chinese_Complete_Workbook_Translation_and_Essay_Edition.docx"
OUTPUT = ROOT / "output" / "Grade_10_Chinese_Student_Practice_Workbook_No_Answers.docx"


def element_text(element):
    return "".join(element.itertext())


doc = Document(SOURCE)
body = doc._element.body
children = list(body.iterchildren())

answer_start = next(
    i for i, element in enumerate(children)
    if "ANSWERS" in element_text(element) and "参考答案" in element_text(element)
)
word_list_start = next(
    i for i, element in enumerate(children[answer_start + 1 :], answer_start + 1)
    if "WORD LIST" in element_text(element) and "总词表" in element_text(element)
)

for element in children[answer_start:word_list_start]:
    body.remove(element)

# Remove the answer-key entry from the contents table.
for table in doc.tables:
    for row in list(table.rows):
        if row.cells and row.cells[0].text.strip() == "Answers":
            table._tbl.remove(row._tr)

# Make the cover clearly identify the classroom handout version.
old_cover_line = "Complete Chinese Activity Workbook - Translation & Essay Edition"
new_cover_line = "Complete Chinese Activity Workbook - Student Practice Edition"
for paragraph in doc.paragraphs:
    if old_cover_line in paragraph.text:
        for index, run in enumerate(paragraph.runs):
            run.text = new_cover_line if index == 0 else ""

doc.core_properties.title = "Grade 10 Chinese Student Practice Workbook - No Answers"
doc.core_properties.subject = "Classroom student edition with blank translation and essay response spaces"
doc.core_properties.comments = "Student practice copy. Answer keys and model translations are intentionally omitted."
doc.save(OUTPUT)
print(OUTPUT)

