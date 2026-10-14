import json
import unittest
from datetime import date
from io import BytesIO
from urllib.parse import urlencode

from openpyxl import load_workbook
import test_app
from app import Category, Group, Transaction, Adjustment, db, build_report, particular_codes


class GroupReportTests(unittest.TestCase):
    login = test_app.FinanceAppTests.login
    csrf = test_app.FinanceAppTests.csrf

    def setUp(self):
        test_app.FinanceAppTests.setUp(self)
        self.login()
        self.day = date(2026, 9, 12)
        self.keys = [json.dumps([kind, 'Custom group']) for kind in ('income', 'expense')]
        with self.app.app_context():
            for kind, amount in [('income', 115), ('expense', 40)]:
                db.session.add(Group(kind=kind, name='Custom group', active=False))
                cat = Category(kind=kind, group_name='Custom group', name=kind + ' account', active=False)
                db.session.add(cat)
                db.session.flush()
                txn = Transaction(category_id=cat.id, cash_date=self.day, reference=kind,
                                  counterparty='Test', subtotal=amount - 15, sscl=5, vat=10,
                                  total=amount, created_by_id=1)
                db.session.add(txn)
                db.session.flush()
                if kind == 'income':
                    db.session.add(Adjustment(transaction_id=txn.id, effective_date=self.day,
                                              amount=-5, reason='Correction', created_by_id=1))
            db.session.add(Category(kind='income', group_name='Custom group', name='Zero account'))
            other = Category.query.filter_by(kind='income', name='Security Service Income').first()
            db.session.add(Transaction(category_id=other.id, cash_date=self.day, reference='excluded',
                                       counterparty='Other', subtotal=800, total=800, created_by_id=1))
            db.session.commit()

    def tearDown(self):
        with self.app.app_context():
            db.session.remove()
            db.engine.dispose()

    def query(self, groups=None, report_type='monthly'):
        return urlencode({'start': str(self.day), 'end': str(self.day), 'type': report_type,
                          'group': self.keys if groups is None else groups}, doseq=True)

    def test_multi_group_totals_and_duplicate_codes(self):
        with self.app.app_context():
            report = build_report(self.day, self.day, selected_groups=self.keys)
            self.assertEqual((report['income'], report['expense'], report['pl']), (110, 40, 70))
            items = report['sections'][0]['groups'][0]['items']
            self.assertEqual(len(items), 1)
            self.assertEqual(sum(i['amounts']['Cash Total'] for i in items), 110)
            single = build_report(self.day, self.day, selected_groups=[self.keys[0]])
            self.assertEqual(single['expense'], 0)
            self.assertFalse(single['sections'][1]['groups'])
            empty = build_report(date(2020, 1, 1), date(2020, 1, 1), selected_groups=self.keys)
            self.assertEqual(empty['pl'], 0)
            self.assertFalse(empty['sections'][0]['groups'])

    def test_html_and_export_share_selection(self):
        response = self.client.get('/reports?' + self.query())
        self.assertEqual(response.status_code, 200)
        self.assertNotIn(b'Zero account', response.data)
        self.assertNotIn(b'excluded', response.data)
        export = self.client.get('/reports/export?' + self.query())
        self.assertEqual(export.status_code, 200)
        wb = load_workbook(BytesIO(export.data))
        self.assertEqual(wb.sheetnames, ['Group Report', 'Summary', 'Detailed Ledger'])
        rows = list(wb['Group Report'].values)
        profit_row = next(row for row in rows if row[1] == 'Liquid P/L')
        self.assertEqual(profit_row[-1], 70)
        self.assertEqual(rows[4], ('No.', 'Particular', 'A/C code', 'Amount (LKR)', 'Total (LKR)'))
        labels = [row[1] for row in rows]
        self.assertLess(labels.index('Total Income'), labels.index('Expenditure Particulars'))
        self.assertLess(labels.index('Total Expenditure'), labels.index('Liquid P/L'))
        with self.app.app_context():
            cat = Category.query.filter_by(kind='income', name='income account').one()
            income_code = f"IP_{particular_codes()[cat.id]:04d}"
        self.assertIn((1, 'income account', income_code, 110, None), rows)
        html = response.data.decode()
        self.assertLess(html.index('Income Particulars'), html.index('Expenditure Particulars'))
        self.assertLess(html.index('class="report-strip"'), html.index('Income Particulars'))
        self.assertIn('class="group-picker" open', html)
        self.assertIn('Custom group', rows[3][0])
        self.assertNotIn('excluded', str(list(wb['Detailed Ledger'].values)))
        short = self.client.get('/reports/export?' + self.query(report_type='short'))
        wb = load_workbook(BytesIO(short.data))
        self.assertNotIn('Detailed Ledger', wb.sheetnames)
        self.assertNotIn('Zero account', str(list(wb['Group Report'].values)))
        short_rows = list(wb['Group Report'].values)
        self.assertEqual(sum(row[0] == 'Custom group' for row in short_rows), 2)
        self.assertIn(('Liquid P/L', 70), short_rows)
        short_html = self.client.get('/reports?' + self.query(report_type='short')).data
        self.assertNotIn(b'income account', short_html)
        self.assertNotIn(b'Custom group \xe2\x80\x94 Total', short_html)
        self.assertNotIn(b'<option value="direct"', response.data)

    def test_invalid_group_and_report_types(self):
        for path in ['/reports', '/reports/export']:
            self.assertEqual(self.client.get(path + '?' + self.query(['invalid'])).status_code, 400)
        for report_type in ['monthly', 'short', 'direct', 'rental', 'construction']:
            self.assertEqual(self.client.get('/reports?' + self.query(report_type=report_type)).status_code, 200)
        self.assertEqual(self.client.get('/reports?' + self.query([])).status_code, 200)

    def test_report_code_order_and_compact_export_styles(self):
        names = ['Trade Income', 'Bank Interest Income', 'Rental Income', 'Construction Projects Income', 'Inward Taxes']
        with self.app.app_context():
            for name in reversed(names):
                cat = Category.query.filter_by(kind='income', group_name=name).first()
                db.session.add(Transaction(category_id=cat.id, cash_date=self.day, reference='ORDER-CHECK',
                    counterparty='', subtotal=100, total=100, created_by_id=1))
            db.session.commit()
            keys = [json.dumps(['income', name]) for name in reversed(names)]
            report = build_report(self.day, self.day, selected_groups=keys)
            self.assertEqual([g['name'] for g in report['sections'][0]['groups']], names)
            self.assertEqual([g['label'] for g in report['summary']], names)
            tax_id = Group.query.filter_by(kind='income', name='Inward Taxes').one().id
        for kind in ['monthly', 'short']:
            response = self.client.get('/reports/export?' + self.query(keys, kind))
            wb = load_workbook(BytesIO(response.data))
            ws = wb['Group Report']
            label_col = 1 if kind == 'short' else 2
            labels = [ws.cell(r, label_col).value for r in range(1, ws.max_row + 1)]
            self.assertEqual(sorted(names, key=labels.index), names)
            section_row = labels.index('Income Particulars') + 1
            heading = ws.cell(section_row, label_col)
            self.assertEqual(heading.fill.fgColor.rgb, '004B6575')
            self.assertEqual(heading.font.sz, 13)
            self.assertEqual(heading.font.color.rgb, '00FFFFFF')
            total_row = labels.index('Total Income') + 1
            for cell in ws[total_row]:
                self.assertEqual(cell.border.top.style, 'thin')
                self.assertEqual(cell.border.bottom.style, 'double')
            self.assertEqual(ws.row_dimensions[section_row].height, 26)
            if kind == 'monthly':
                body_row = next(r for r in range(6, ws.max_row + 1) if ws.cell(r, 3).value)
                self.assertEqual(ws.row_dimensions[body_row].height, 20)
                self.assertEqual(ws.cell(body_row, 2).font.sz, 11)
        # The report follows administrator-edited numeric codes, not alphabetical names.
        self.client.post(f'/groups/{tax_id}/edit', data={'csrf_token':self.csrf(), 'name':'Inward Taxes', 'code':'5'})
        with self.app.app_context():
            report = build_report(self.day, self.day, selected_groups=keys)
            self.assertEqual(report['sections'][0]['groups'][0]['name'], 'Inward Taxes')
